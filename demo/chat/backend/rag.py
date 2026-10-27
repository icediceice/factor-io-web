"""Parsing, chunking and retrieval for the Factor IO chat demo.

Everything here lives in memory for one gateway process on the Colab VM. A
customer's documents never touch disk beyond the upload buffer, and they are
gone when the VM stops: that is the privacy story the demo page tells, so do
not add persistence here without changing the page.

Retrieval is hybrid: BM25 for exact terms (part numbers, names, figures) and a
small multilingual embedding model for paraphrase, fused by reciprocal rank.
Thai and CJK have no spaces, so BM25 indexes their runs as character bigrams.
"""
from __future__ import annotations

import csv
import hashlib
import html.parser
import io
import math
import re
import uuid
from collections import Counter
from dataclasses import dataclass, field

import numpy as np

# ~3 characters per token is a deliberate over-estimate for English and close
# for Thai; over-estimating only means retrieval kicks in a little earlier.
CHARS_PER_TOKEN = 3
CHUNK_CHARS = 2400
CHUNK_OVERLAP = 400
MAX_SHEET_ROWS = 5000

TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".log", ".xml", ".yaml", ".yml", ".rst"}
SUPPORTED_EXT = TEXT_EXT | {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".tsv", ".html", ".htm"}


class UnsupportedFile(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


@dataclass
class Section:
    text: str
    ref: str = ""  # "p. 3", "sheet Sales", "" for flat text


@dataclass
class Chunk:
    id: int
    doc_id: str
    doc_name: str
    ref: str
    text: str


@dataclass
class Document:
    id: str
    name: str
    chars: int
    tokens: int
    chunk_ids: list[int] = field(default_factory=list)

    def summary(self) -> dict:
        return {"id": self.id, "name": self.name, "tokens": self.tokens, "chunks": len(self.chunk_ids)}


def estimate_tokens(text: str) -> int:
    return max(1, len(text) // CHARS_PER_TOKEN)


# ---------------------------------------------------------------- extraction

def _ext(name: str) -> str:
    dot = name.rfind(".")
    return name[dot:].lower() if dot >= 0 else ""


def decode_text(data: bytes) -> str:
    # utf-8-sig first; cp874 is the Thai Windows code page, common in exported CSVs.
    for enc in ("utf-8-sig", "cp874"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


class _HTMLText(html.parser.HTMLParser):
    SKIP = {"script", "style", "noscript", "template"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self):
        super().__init__()
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def _rows_to_text(rows) -> str:
    lines = []
    for row in rows:
        cells = ["" if c is None else str(c).strip() for c in row]
        if any(cells):
            lines.append(" | ".join(cells))
        if len(lines) >= MAX_SHEET_ROWS:
            lines.append(f"... (truncated at {MAX_SHEET_ROWS} rows)")
            break
    return "\n".join(lines)


def extract(name: str, data: bytes) -> list[Section]:
    ext = _ext(name)
    if ext not in SUPPORTED_EXT:
        raise UnsupportedFile(f"{name}: unsupported type {ext or '(none)'}")

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        sections = [Section(page.extract_text() or "", f"p. {i}") for i, page in enumerate(reader.pages, 1)]
    elif ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts.append(_rows_to_text([c.text for c in r.cells] for r in t.rows))
        sections = [Section("\n\n".join(parts))]
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections = [Section(_rows_to_text(ws.iter_rows(values_only=True)), f"sheet {ws.title}") for ws in wb.worksheets]
    elif ext in (".csv", ".tsv"):
        text = decode_text(data)
        delim = "\t" if ext == ".tsv" else ","
        if ext == ".csv":
            try:
                delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
            except csv.Error:
                pass
        sections = [Section(_rows_to_text(csv.reader(io.StringIO(text), delimiter=delim)))]
    elif ext in (".html", ".htm"):
        p = _HTMLText()
        p.feed(decode_text(data))
        sections = [Section("".join(p.out))]
    else:
        sections = [Section(decode_text(data))]

    sections = [Section(_normalise(s.text), s.ref) for s in sections]
    sections = [s for s in sections if s.text]
    if not sections:
        raise EmptyDocument(f"{name}: no readable text (a scanned PDF needs OCR, which this demo does not run)")
    return sections


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ------------------------------------------------------------------ chunking

def _hard_split(text: str, size: int) -> list[str]:
    out = []
    while len(text) > size:
        cut = max(text.rfind(". ", 0, size), text.rfind("\n", 0, size), text.rfind(" ", 0, size))
        cut = cut + 1 if cut > size // 2 else size
        out.append(text[:cut].strip())
        text = text[cut:]
    if text.strip():
        out.append(text.strip())
    return out


def chunk_section(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Pack paragraphs into ~size-char chunks; each chunk after the first starts
    with the tail of the previous one so a fact split across a boundary survives."""
    paras = []
    for p in text.split("\n\n"):
        paras.extend(_hard_split(p, size) if len(p) > size else [p])
    chunks, cur = [], ""
    for p in paras:
        if cur and len(cur) + 2 + len(p) > size:
            chunks.append(cur)
            tail = cur[-overlap:]
            sp = tail.find(" ")
            cur = (tail[sp + 1:] if 0 <= sp < overlap // 2 else tail) + "\n\n" + p
        else:
            cur = f"{cur}\n\n{p}" if cur else p
    if cur.strip():
        chunks.append(cur)
    return chunks


# ----------------------------------------------------------------- tokenising

_TOKEN_RE = re.compile(r"[0-9a-zÀ-ɏ]+|[฀-๿]+|[぀-ヿ一-鿿가-힯]+")
_ASCII_WORD = re.compile(r"[0-9a-zÀ-ɏ]+")


def tokenize(text: str) -> list[str]:
    toks = []
    for run in _TOKEN_RE.findall(text.lower()):
        if _ASCII_WORD.fullmatch(run):
            toks.append(run)
        elif len(run) == 1:
            toks.append(run)
        else:
            toks.extend(run[i:i + 2] for i in range(len(run) - 1))
    return toks


class BM25:
    def __init__(self, k1: float = 1.5, b: float = 0.75):
        self.k1, self.b = k1, b
        self.docs: dict[int, Counter] = {}
        self.lens: dict[int, int] = {}
        self.df: Counter = Counter()

    def add(self, cid: int, text: str):
        tf = Counter(tokenize(text))
        self.docs[cid], self.lens[cid] = tf, sum(tf.values())
        self.df.update(tf.keys())

    def remove(self, cid: int):
        tf = self.docs.pop(cid, None)
        if tf is not None:
            self.lens.pop(cid)
            self.df.subtract(tf.keys())
            self.df += Counter()  # drop zero counts

    def scores(self, query: str) -> dict[int, float]:
        n = len(self.docs)
        if not n:
            return {}
        avg = sum(self.lens.values()) / n
        out: dict[int, float] = {}
        for term in set(tokenize(query)):
            df = self.df.get(term, 0)
            if not df:
                continue
            idf = math.log(1 + (n - df + 0.5) / (df + 0.5))
            for cid, tf in self.docs.items():
                f = tf.get(term)
                if f:
                    denom = f + self.k1 * (1 - self.b + self.b * self.lens[cid] / avg)
                    out[cid] = out.get(cid, 0.0) + idf * f * (self.k1 + 1) / denom
        return out


# ----------------------------------------------------------------- embedders

class HashEmbedder:
    """Deterministic bag-of-hashed-tokens vectors: no model download. Used by
    the tests and as a fallback when the real model cannot load."""
    name = "hash-256"

    def __init__(self, dim: int = 256):
        self.dim = dim

    def _embed(self, texts: list[str]) -> np.ndarray:
        out = np.zeros((len(texts), self.dim), dtype=np.float32)
        for i, t in enumerate(texts):
            for tok in tokenize(t):
                h = int.from_bytes(hashlib.blake2b(tok.encode(), digest_size=4).digest(), "little")
                out[i, h % self.dim] += 1.0
        norms = np.linalg.norm(out, axis=1, keepdims=True)
        return out / np.where(norms == 0, 1, norms)

    def passages(self, texts: list[str]) -> np.ndarray:
        return self._embed(texts)

    def query(self, text: str) -> np.ndarray:
        return self._embed([text])[0]


class E5Embedder:
    """intfloat/multilingual-e5-small (118M params) on CPU: covers Thai and
    English, and leaves the whole GPU to Strata's expert cache."""
    name = "multilingual-e5-small"

    def __init__(self, model: str = "intfloat/multilingual-e5-small", device: str = "cpu"):
        from sentence_transformers import SentenceTransformer
        self.model = SentenceTransformer(model, device=device)

    def passages(self, texts: list[str]) -> np.ndarray:
        return self.model.encode([f"passage: {t}" for t in texts], batch_size=32,
                                 normalize_embeddings=True, convert_to_numpy=True).astype(np.float32)

    def query(self, text: str) -> np.ndarray:
        return self.model.encode([f"query: {text}"], normalize_embeddings=True,
                                 convert_to_numpy=True)[0].astype(np.float32)


def get_embedder(kind: str):
    if kind == "hash":
        return HashEmbedder()
    try:
        return E5Embedder()
    except Exception as exc:  # model download or import failure must not take the demo down
        print(f"[rag] embedding model unavailable ({type(exc).__name__}: {exc}); falling back to hash vectors", flush=True)
        return HashEmbedder()


# --------------------------------------------------------------------- store

SYSTEM_PROMPT = """You are a document assistant in a live product demo. The user uploaded the documents excerpted below.

Rules:
- Answer only from the numbered sources. Cite every claim with its source number in square brackets, like [2] or [1][4].
- If the sources do not contain the answer, say so plainly and say what the documents do cover. Never invent figures, names or dates.
- For tables, quote the exact values and say which row or sheet they come from.
- Reply in the language of the user's question.
- Be concise: lead with the answer, then the supporting detail."""

NO_DOCS_PROMPT = """You are a helpful assistant in a live product demo. No documents have been uploaded yet.
Answer briefly, and if the question needs the user's own data, suggest they upload the relevant file with the + button. Reply in the language of the user's question."""


class DocStore:
    """One customer session's documents and indexes."""

    def __init__(self, embedder):
        self.embedder = embedder
        self.docs: dict[str, Document] = {}
        self.chunks: dict[int, Chunk] = {}
        self.vectors: dict[int, np.ndarray] = {}
        self.bm25 = BM25()
        self._next = 0

    @property
    def total_tokens(self) -> int:
        return sum(d.tokens for d in self.docs.values())

    def add_document(self, name: str, data: bytes) -> Document:
        sections = extract(name, data)
        doc = Document(id=uuid.uuid4().hex[:12], name=name,
                       chars=sum(len(s.text) for s in sections), tokens=0)
        new: list[Chunk] = []
        for s in sections:
            for text in chunk_section(s.text):
                new.append(Chunk(self._next, doc.id, name, s.ref, text))
                self._next += 1
        vecs = self.embedder.passages([c.text for c in new])
        for c, v in zip(new, vecs):
            self.chunks[c.id], self.vectors[c.id] = c, v
            self.bm25.add(c.id, c.text)
            doc.chunk_ids.append(c.id)
        doc.tokens = estimate_tokens("".join(s.text for s in sections))
        self.docs[doc.id] = doc
        return doc

    def remove_document(self, doc_id: str) -> bool:
        doc = self.docs.pop(doc_id, None)
        if doc is None:
            return False
        for cid in doc.chunk_ids:
            self.chunks.pop(cid, None)
            self.vectors.pop(cid, None)
            self.bm25.remove(cid)
        return True

    def search(self, query: str, k: int) -> list[Chunk]:
        if not self.chunks:
            return []
        ids = list(self.chunks)
        q = self.embedder.query(query)
        dense = np.stack([self.vectors[i] for i in ids]) @ q
        dense_rank = [ids[i] for i in np.argsort(-dense)]
        sparse = self.bm25.scores(query)
        sparse_rank = sorted(sparse, key=sparse.get, reverse=True)
        fused: dict[int, float] = {}
        for ranking in (dense_rank, sparse_rank):
            for r, cid in enumerate(ranking):
                fused[cid] = fused.get(cid, 0.0) + 1.0 / (60 + r)
        best = sorted(fused, key=fused.get, reverse=True)[:k]
        return [self.chunks[c] for c in best]

    def select(self, query: str, stuff_tokens: int, top_k: int, budget_tokens: int) -> tuple[str, list[Chunk]]:
        """Whole documents when they fit stuff_tokens, else the top_k fused hits
        trimmed to budget_tokens. Returns (mode, chunks in reading order)."""
        if not self.chunks:
            return "none", []
        if self.total_tokens <= stuff_tokens:
            return "full", [self.chunks[c] for c in sorted(self.chunks)]
        picked, used = [], 0
        for c in self.search(query, top_k):
            t = estimate_tokens(c.text)
            if picked and used + t > budget_tokens:
                break
            picked.append(c)
            used += t
        return "retrieval", sorted(picked, key=lambda c: c.id)


def build_messages(chunks: list[Chunk], history: list[dict], question: str) -> tuple[list[dict], list[dict]]:
    """Returns (OpenAI messages, citations). Source n in the prompt is citations[n-1]."""
    citations = []
    if chunks:
        blocks = []
        for n, c in enumerate(chunks, 1):
            where = f"{c.doc_name}, {c.ref}" if c.ref else c.doc_name
            blocks.append(f"[{n}] ({where})\n{c.text}")
            citations.append({"n": n, "doc": c.doc_name, "doc_id": c.doc_id, "ref": c.ref, "text": c.text})
        system = SYSTEM_PROMPT + "\n\nSOURCES\n\n" + "\n\n---\n\n".join(blocks)
    else:
        system = NO_DOCS_PROMPT
    messages = [{"role": "system", "content": system}]
    for turn in history:
        if turn.get("role") in ("user", "assistant") and isinstance(turn.get("content"), str):
            messages.append({"role": turn["role"], "content": turn["content"]})
    messages.append({"role": "user", "content": question})
    return messages, citations